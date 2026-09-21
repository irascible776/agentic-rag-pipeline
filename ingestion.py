# ingestion.py - Document Loading, Chunking, and Dual-Mode Ingestion (Supabase & Local Chroma)
import os
import sys
from pathlib import Path
from typing import List, Tuple, Dict, Any, Union
from pydantic import SecretStr

import io
import pypdf
import docx

# Add root directory to sys.path
root_dir = Path(__file__).resolve().parent.parent
if str(root_dir) not in sys.path:
    sys.path.insert(0, str(root_dir))

import config

from langchain_core.documents import Document
from langchain_text_splitters.character import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings


def get_embedding_model():
    """Returns Gemini embedding model (runs on Google Cloud, uses 0 MB server RAM)."""
    return GoogleGenerativeAIEmbeddings(
        model=config.EMBEDDING_MODEL,
        api_key=SecretStr(config.GEMINI_API_KEY) if config.GEMINI_API_KEY else None,
        output_dimensionality=config.EMBEDDING_DIM
    )


def extract_text_from_bytes(filename: str, raw_bytes: bytes) -> str:
    # get filename in lowercase to check extension
    name_lower = filename.lower()

    # 1. extract text from pdf
    if name_lower.endswith(".pdf"):
        reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
        pages_text: List[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(text.strip())
        return "\n\n".join(pages_text)

    # 2. extract text from word docx
    elif name_lower.endswith(".docx"):
        doc = docx.Document(io.BytesIO(raw_bytes))
        paragraphs: List[str] = []
        # get paragraph text
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())
        # also read text inside tables
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    paragraphs.append(row_str)
        return "\n\n".join(paragraphs)

    # 3. extract text from plain txt or markdown
    else:
        return raw_bytes.decode("utf-8", errors="ignore").strip()


def load_docs_from_folder(docs_path: str | None = None) -> List[Document]:
    # load .txt, .pdf, .docx, and .md files from a local folder
    folder = Path(docs_path or config.LOCAL_DOCS_DIR)
    print(f"[Ingestion] Loading documents from folder: {folder}...")

    documents: List[Document] = []
    if not folder.exists():
        print(f"[Ingestion] Directory '{folder}' does not exist.")
        return []

    # supported document formats
    supported_exts = {".txt", ".md", ".pdf", ".docx"}

    for file_path in sorted(folder.iterdir()):
        if file_path.suffix.lower() not in supported_exts:
            continue
        try:
            # read file bytes and extract text
            with open(file_path, "rb") as f:
                content = extract_text_from_bytes(file_path.name, f.read())
            if content.strip():
                documents.append(Document(
                    page_content=content.strip(),
                    metadata={"source": file_path.name}
                ))
        except Exception as e:
            print(f"[Ingestion] Error reading {file_path.name}: {e}")

    print(f"[Ingestion] Successfully loaded {len(documents)} document(s).")
    return documents


def load_docs_from_memory(files: List[Tuple[str, str]]) -> List[Document]:
    """
    Loads documents from in-memory tuples: [(filename, text_content), ...].
    Used by the Web GUI when users upload files directly in the browser.
    """
    documents = []
    for filename, content in files:
        if content and content.strip():
            doc = Document(
                page_content=content.strip(),
                metadata={"source": filename}
            )
            documents.append(doc)
    print(f"[Ingestion] Loaded {len(documents)} document(s) from upload payload.")
    return documents


def chunk_docs(docs: List[Document], size: int | None = None, overlap: int | None = None) -> List[Document]:
    """Splits documents into smaller overlapping text chunks for semantic retrieval."""
    chunk_size = size or config.CHUNK_SIZE
    chunk_overlap = overlap or config.CHUNK_OVERLAP

    print(f"[Ingestion] Chunking {len(docs)} document(s) (size={chunk_size}, overlap={chunk_overlap})...")

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=config.CHUNK_SEPARATORS
    )

    chunks = text_splitter.split_documents(docs)
    print(f"[Ingestion] Created {len(chunks)} text chunk(s).")
    return chunks


def store_chunks(chunks: List[Document], session_id: str = "global", is_permanent: bool = False) -> Dict[str, Any]:
    """
    Stores chunks in the vector database:
    - Primary: Supabase pgvector (Cloud) if SUPABASE_URL is set
    - Fallback: Local ChromaDB if SUPABASE_URL is missing
    """
    target_session = "global" if is_permanent else session_id
    embed_model = get_embedding_model()

    # --- MODE 1: CLOUD SUPABASE ---
    if config.USE_SUPABASE:
        print(f"[Ingestion] Cloud Mode: Indexing {len(chunks)} chunks into Supabase table '{config.SUPABASE_TABLE}'...")
        from supabase import create_client
        supabase = create_client(config.SUPABASE_URL, config.SUPABASE_KEY)

        # 1. Compute embeddings for all chunk texts
        texts = [chunk.page_content for chunk in chunks]
        embeddings = embed_model.embed_documents(texts)

        # 2. Build rows with session_id tag
        records = []
        for chunk, vec in zip(chunks, embeddings):
            meta = dict(chunk.metadata)
            meta["session_id"] = target_session
            records.append({
                "session_id": target_session,
                "content": chunk.page_content,
                "metadata": meta,
                "embedding": vec
            })

        # 3. Insert records in batches of 50
        batch_size = 50
        for i in range(0, len(records), batch_size):
            batch = records[i:i + batch_size]
            supabase.table(config.SUPABASE_TABLE).insert(batch).execute()

        print(f"[Ingestion] Successfully stored {len(chunks)} chunks in Supabase under session '{target_session}'.")
        return {"mode": "supabase", "chunks_stored": len(chunks), "session_id": target_session}

    # --- MODE 2: LOCAL CHROMADB FALLBACK ---
    else:
        print(f"[Ingestion] Local Mode: Indexing {len(chunks)} chunks into Chroma at {config.LOCAL_DB_DIR}...")
        # Lazy import of Chroma so it only loads when local mode is active
        from langchain_chroma import Chroma

        # Tag metadata with session_id
        for chunk in chunks:
            chunk.metadata["session_id"] = target_session

        # Save to local Chroma store
        vectordb = Chroma.from_documents(
            documents=chunks,
            embedding=embed_model,
            persist_directory=str(config.LOCAL_DB_DIR),
            collection_name="gemini_rag",
            collection_metadata={"hnsw:space": "cosine"}
        )

        print(f"[Ingestion] Successfully stored {len(chunks)} chunks in local Chroma under session '{target_session}'.")
        return {"mode": "local", "chunks_stored": len(chunks), "session_id": target_session}


def ingest_documents(
    files_or_folder: Union[str, List[Tuple[str, str]], None] = None,
    session_id: str = "global",
    is_permanent: bool = False,
    chunk_size: int | None = None,
    chunk_overlap: int | None = None
) -> Dict[str, Any]:
    """
    Standard ingestion pipeline:
    1. Loads text files (from folder path or memory list of (filename, text))
    2. Chunks them (using optional custom chunk_size and chunk_overlap)
    3. Embeds & stores in Supabase or local Chroma
    """
    if files_or_folder is None:
        files_or_folder = str(config.LOCAL_DOCS_DIR)

    if isinstance(files_or_folder, str):
        docs = load_docs_from_folder(files_or_folder)
    elif isinstance(files_or_folder, list):
        docs = load_docs_from_memory(files_or_folder)
    else:
        raise ValueError("files_or_folder must be a folder path string or a list of (filename, content) tuples.")

    if not docs:
        return {"status": "error", "message": "No documents found to ingest.", "chunks_count": 0, "files_count": 0}

    chunks = chunk_docs(docs, size=chunk_size, overlap=chunk_overlap)
    result = store_chunks(chunks, session_id=session_id, is_permanent=is_permanent)

    return {
        "status": "success",
        "files_count": len(docs),
        "chunks_count": len(chunks),
        "mode": result["mode"],
        "session_id": result["session_id"]
    }


def main():
    """CLI test runner for ingestion."""
    print("=" * 50)
    print("  Running Ingestion Pipeline")
    print("=" * 50)
    result = ingest_documents(str(config.LOCAL_DOCS_DIR), session_id="global", is_permanent=True)
    print("\nIngestion Result:", result)


if __name__ == "__main__":
    main()
